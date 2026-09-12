from dataclasses import dataclass
from datetime import date, datetime
import math
import os
from typing import BinaryIO
import unicodedata
import zipfile
import openpyxl
from openpyxl.utils.exceptions import InvalidFileException
from soldoni.core.models import OpType, Transaction

HEADER_ROW = 5      # indice 0-based della riga header
FIRST_DATA_ROW = 7  # indice 0-based della prima riga dati
COMMISSION_COLS = (11, 12, 13, 14)
_EXPECTED_HEADERS = {
    0: "operazione", 1: "data valuta", 2: "descrizione", 3: "titolo",
    4: "isin", 5: "segno", 6: "quantita", 7: "divisa", 8: "prezzo",
    9: "cambio", 10: "controvalore",
}


@dataclass(frozen=True)
class ImportedInstrument:
    isin: str
    name: str
    native_currency: str


def _to_date(value) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return datetime.strptime(str(value).strip(), "%d/%m/%Y").date()


def _num(value, field: str = "valore") -> float:
    if value in (None, ""):
        return 0.0
    number = float(value)
    if not math.isfinite(number):
        raise ValueError(f"{field} non è un numero finito")
    return number


def _normalized_header(value) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    return " ".join("".join(c for c in text if not unicodedata.combining(c)).lower().split())


def _validate_headers(rows) -> None:
    if len(rows) <= HEADER_ROW:
        raise ValueError("Intestazione Fineco non trovata")
    header = rows[HEADER_ROW]
    if len(header) <= max(_EXPECTED_HEADERS):
        raise ValueError("Intestazione Fineco incompleta")
    invalid = [
        expected
        for index, expected in _EXPECTED_HEADERS.items()
        if _normalized_header(header[index]) != expected
    ]
    if invalid:
        raise ValueError("Colonne Fineco mancanti o non riconosciute: " + ", ".join(invalid))


def _op_type(desc: str, segno: str) -> OpType:
    desc = (desc or "").strip()
    segno = (segno or "").strip()
    if desc == "Dividendo":
        return OpType.DIVIDEND
    if desc == "Stacco Cedole":
        return OpType.DIVIDEND
    if desc == "Rimborso":
        return OpType.SELL
    if desc == "Compravendita titoli":
        if segno == "A":
            return OpType.BUY
        if segno == "V":
            return OpType.SELL
        raise ValueError(f"Segno non riconosciuto '{segno}' per compravendita")
    raise ValueError(f"Descrizione operazione non gestita: '{desc}'")


def parse_fineco_export(path: str) -> list[Transaction]:
    transactions, _ = parse_fineco_export_with_instruments(path)
    return transactions


def parse_fineco_export_with_instruments(
        path: str | os.PathLike | BinaryIO,
) -> tuple[list[Transaction], dict[str, ImportedInstrument]]:
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        try:
            ws = (wb["Movimenti Dossier Titoli"]
                  if "Movimenti Dossier Titoli" in wb.sheetnames else wb.active)
            rows = list(ws.iter_rows(values_only=True))
        finally:
            wb.close()
    except (OSError, KeyError, TypeError, ValueError, zipfile.BadZipFile,
            InvalidFileException) as e:
        raise ValueError(f"File Excel non leggibile: {e}") from e
    _validate_headers(rows)
    out: list[Transaction] = []
    instruments: dict[str, ImportedInstrument] = {}
    for row_index, raw in enumerate(rows[FIRST_DATA_ROW:], start=FIRST_DATA_ROW + 1):
        if raw is None or all(c in (None, "") for c in raw):
            continue
        try:
            if len(raw) <= max(COMMISSION_COLS):
                raise ValueError("riga incompleta")
            desc = str(raw[2] or "").strip()
            if not desc:
                raise ValueError("descrizione mancante")
            op = _op_type(desc, raw[5])
            commission = sum(_num(raw[c], "commissione") for c in COMMISSION_COLS)
            isin = str(raw[4] or "").strip()
            if not isin:
                raise ValueError("ISIN mancante")
            trade_date = _to_date(raw[0])
            value_date = _to_date(raw[1])
            quantity = _num(raw[6], "quantità")
            amount_eur = _num(raw[10], "controvalore")
            price_native = _num(raw[8], "prezzo")
            fx_rate = _num(raw[9], "cambio") or 1.0
        except (IndexError, TypeError, ValueError) as e:
            raise ValueError(f"Riga Excel {row_index}: {e}") from e
        if isin not in instruments:
            instruments[isin] = ImportedInstrument(
                isin=isin,
                name=str(raw[3] or isin).strip(),
                native_currency=str(raw[7] or "EUR").strip().upper(),
            )
        out.append(Transaction(
            trade_date=trade_date,
            value_date=value_date,
            isin=isin,
            op_type=op,
            quantity=quantity,
            amount_eur=amount_eur,
            price_native=price_native,
            fx_rate=fx_rate,
            commission_eur=commission,
            source="import",
            raw_desc=str(desc).strip(),
        ))
    return out, instruments
